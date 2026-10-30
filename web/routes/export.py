from io import BytesIO

from flask import Blueprint, send_file
from openpyxl import Workbook
from openpyxl.styles import Font

from common.db import connect, load_programs
from web.auth import require_session
from web.views.letters import all_letter_rows

export_bp = Blueprint("export", __name__, url_prefix="/api/v1")

ANALYSIS_STATUS_LABELS = {
    "discovered": "ממתין לניתוח",
    "queued": "בתור לניתוח",
    "analyzing": "הניתוח מתבצע",
    "analyzed": "נותח",
    "failed": "הניתוח נכשל",
}

HEADERS = [
    ("letter_date", "תאריך"),
    ("request_numbers", "מספר בקשה"),
    ("ministry", "משרד"),
    ("analysis_status", "סטטוס ניתוח"),
    ("relevant", "רלוונטי"),
    ("requires_review", "דורש עיון"),
    ("passed", "עברה"),
    ("last_sent_at", "נשלח"),
    ("send_count", "מספר שליחות"),
    ("letter_url", "קישור PDF"),
]


@export_bp.get("/letters/export")
@require_session
def export_letters():
    with connect() as conn:
        programs = list(load_programs(conn))
        rows = all_letter_rows(conn, programs=programs)
    wb = Workbook()
    ws = wb.active
    ws.title = "פניות"
    ws.append([label for _, label in HEADERS])
    link_font = Font(color="0563C1", underline="single")
    for item in rows:
        line = []
        for key, _ in HEADERS:
            val = item.get(key)
            if key == "analysis_status":
                val = ANALYSIS_STATUS_LABELS.get(val, val)
            elif key == "request_numbers" and isinstance(val, list):
                val = ", ".join(val)
            elif isinstance(val, bool):
                val = "כן" if val else "לא"
            line.append(val)
        ws.append(line)
        url = item.get("letter_url")
        if url:
            cell = ws.cell(row=ws.max_row, column=len(HEADERS))
            cell.hyperlink = url
            cell.font = link_font
            cell.value = url
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="letters.xlsx")
