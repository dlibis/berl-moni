"""Program master .xlsx (sheet תוכניות, columns קוד / שם תוכנית). openpyxl only."""
import re
from io import BytesIO

from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook

SHEET = "תוכניות"
CODE_COL = "קוד"
NAME_COL = "שם תוכנית"
CODE_RE = re.compile(r"^[0-9]{1,6}$")


def _norm_code(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not CODE_RE.fullmatch(text):
        return None
    return text.zfill(6)


def programs_to_xlsx(programs: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([CODE_COL, NAME_COL])
    for row in programs:
        ws.append([row["code"], row["name"]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_to_programs(data: bytes) -> list[tuple[str, str]]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        raise ValueError("invalid xlsx") from None
    if SHEET not in wb.sheetnames:
        raise ValueError("missing sheet")
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError("empty sheet")
    cols = {str(c).strip(): i for i, c in enumerate(header) if c is not None}
    if CODE_COL not in cols or NAME_COL not in cols:
        raise ValueError("bad columns")
    ci, ni = cols[CODE_COL], cols[NAME_COL]
    seen: set[str] = set()
    out: list[tuple[str, str]] = []
    for row in rows:
        if row is None:
            continue
        code = _norm_code(row[ci] if ci < len(row) else None)
        name = row[ni] if ni < len(row) else None
        if code is None:
            if name is None or not str(name).strip():
                continue
            raise ValueError("bad code")
        name_s = str(name).strip() if name is not None else ""
        if not name_s:
            raise ValueError("empty name")
        if code in seen:
            raise ValueError("duplicate code")
        seen.add(code)
        out.append((code, name_s))
    if not out:
        raise ValueError("no programs")
    return out
