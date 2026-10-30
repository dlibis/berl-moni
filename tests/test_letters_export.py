from datetime import date
from io import BytesIO

import psycopg
import pytest
from flask import Flask
from openpyxl import load_workbook
from psycopg.rows import dict_row

from tests.test_letters_api import SessionResponse, insert
from web.csrf import register_csrf
from web.routes.export import export_bp


@pytest.fixture
def db(test_db_url):
    connection = lambda: psycopg.connect(test_db_url, row_factory=dict_row)
    with connection() as conn:
        conn.execute("TRUNCATE letters, programs RESTART IDENTITY CASCADE")
        conn.execute("INSERT INTO programs (code, name) VALUES ('200102', 'b')")
    return connection


def _app(db):
    app = Flask(__name__)
    register_csrf(app)
    app.register_blueprint(export_bp)
    return app


@pytest.fixture
def client(db, monkeypatch):
    monkeypatch.setenv("SITE_ORIGIN", "https://berl.site")
    monkeypatch.setenv("NEON_AUTH_BASE_URL", "https://auth.example")
    monkeypatch.setattr("web.auth.requests.get", lambda *args, **kwargs: SessionResponse())
    monkeypatch.setattr("web.routes.export.connect", db)
    return _app(db).test_client()


class UnauthorizedResponse:
    status_code = 401

    def json(self):
        return {}


@pytest.fixture
def anon_client(db, monkeypatch):
    monkeypatch.setenv("SITE_ORIGIN", "https://berl.site")
    monkeypatch.setenv("NEON_AUTH_BASE_URL", "https://auth.example")
    monkeypatch.setattr("web.auth.requests.get", lambda *args, **kwargs: UnauthorizedResponse())
    monkeypatch.setattr("web.routes.export.connect", db)
    return _app(db).test_client()


def test_export_requires_session(anon_client):
    assert anon_client.get("/api/v1/letters/export").status_code == 401


def test_export_xlsx_rows_and_hyperlink(client, db):
    insert(db, 1, day=date(2026, 3, 1))
    insert(db, 2, day=date(2026, 3, 2))
    with db() as conn:
        count = conn.execute("SELECT count(*) AS c FROM letters").fetchone()["c"]
    response = client.get("/api/v1/letters/export", headers={"Origin": "https://berl.site"})
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(BytesIO(response.data))
    ws = wb.active
    assert ws.max_row == count + 1
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "דורש עיון" in headers and "עברה" in headers
    link_cell = ws.cell(row=2, column=len(headers))
    assert link_cell.hyperlink is not None
    assert link_cell.hyperlink.target.endswith(".pdf")


@pytest.mark.parametrize("status,label", [
    ("discovered", "ממתין לניתוח"),
    ("queued", "בתור לניתוח"),
    ("analyzing", "הניתוח מתבצע"),
    ("analyzed", "נותח"),
    ("failed", "הניתוח נכשל"),
])
def test_export_analysis_status_in_hebrew(client, db, status, label):
    insert(db, 1, status=status)
    response = client.get("/api/v1/letters/export")
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.data)).active
    headers = [cell.value for cell in ws[1]]
    assert ws.cell(row=2, column=headers.index("סטטוס ניתוח") + 1).value == label
